"""Read-only workbook validation and transactional, stable-key source import."""
import hashlib
import json
import os
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

import models
from excel_loader import FINANCE_FILE, RD_FILE
from event_service import audit
from governance_service import record_question

UNAVAILABLE = "Mapping unavailable — pending source validation."
SME_REGISTER_NAME = "SME_Clarifications_Register - v 1.1 (1).xlsx"
SME_REGISTER_PATH = Path(os.environ.get("SME_CLARIFICATIONS_FILE", Path.home() / "Downloads" / SME_REGISTER_NAME))


def inspect_sources():
    """Read every occupied cell, including formulas, without exporting a copy."""
    inventory = {}
    for path in (FINANCE_FILE, RD_FILE):
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            inventory[path.name] = {
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                "sheets": [{"name": sheet.title, "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "nonempty_rows": sum(any(c.value is not None for c in row) for row in sheet),
                    "formula_cells": sum(c.data_type == "f" for row in sheet for c in row)}
                    for sheet in workbook],
            }
        finally:
            workbook.close()
    return inventory


def inspect_sme_register():
    result = {"name": SME_REGISTER_NAME, "available": SME_REGISTER_PATH.exists()}
    if not SME_REGISTER_PATH.exists():
        return result
    workbook = load_workbook(SME_REGISTER_PATH, read_only=True, data_only=True)
    try:
        rows = list(workbook["SME Clarifications"].iter_rows(min_row=2, values_only=True))
        responses = {row[0]: row[8] for row in rows if row[0] is not None}
        expected = {1: "difficult", 2: "30%", 3: "Not fail", 4: "all skill", 5: "30% - difficult",
                    7: "keep 20", 13: "70% new", 14: "Objective Score"}
        missing = [key for key, text in expected.items() if text.casefold() not in str(responses.get(key) or "").casefold()]
        result.update(sha256=hashlib.sha256(SME_REGISTER_PATH.read_bytes()).hexdigest(),
                      sheets=workbook.sheetnames, validated=not missing,
                      unresolved_response_ids=missing)
        return result
    finally:
        workbook.close()


def table(path, sheet, header, required):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet]
        columns = [str(c.value).strip() if c.value is not None else None for c in ws[header]]
        if not set(required).issubset(columns):
            raise ValueError(f"{path.name}/{sheet}: required columns missing")
        return [(i, dict(zip(columns, values))) for i, values in
                enumerate(ws.iter_rows(min_row=header + 1, values_only=True), header + 1)
                if any(v is not None for v in values)]
    finally:
        workbook.close()


def source_plan():
    report = {"inventory": inspect_sources(), "sme_register": inspect_sme_register(),
              "issues": [], "counts": {}}
    records = []

    def issue(path, sheet, row, reason, key=None):
        report["issues"].append({"workbook": path.name, "sheet": sheet,
                                 "row": row, "key": key, "reason": reason})

    def unique(path, sheet, header, key, required):
        rows = table(path, sheet, header, [key, *required])
        counts = Counter(str(r.get(key) or "") for _, r in rows)
        result = []
        for n, r in rows:
            value = str(r.get(key) or "")
            if not value or counts[value] != 1:
                issue(path, sheet, n, "Missing or duplicate source key", value)
            else:
                result.append((n, r))
        return result

    def add(path, sheet, row, key, kind, data):
        records.append(dict(workbook=path.name, sheet=sheet, source_row=row,
                            source_key=str(key), entity_type=kind, details=data))

    fskills = unique(FINANCE_FILE, "Skill_Master", 3, "skill_id", ["skill", "capability"])
    rskills = unique(RD_FILE, "Skill Master", 4, "Skill_ID", ["Skill", "Family"])
    names = Counter(str(r.get("skill") or r.get("Skill")) for _, r in fskills + rskills)
    skill_ids = {}
    for path, sheet, rows, key, name, category, description in (
        (FINANCE_FILE, "Skill_Master", fskills, "skill_id", "skill", "capability", "L3_indicator"),
        (RD_FILE, "Skill Master", rskills, "Skill_ID", "Skill", "Family", "Skill definition")):
        for n, r in rows:
            if not r.get(name):
                issue(path, sheet, n, "Missing skill name", r[key]); continue
            # Qualification preserves distinct source IDs where source names repeat.
            display = str(r[name]) if names[str(r[name])] == 1 else f"{r[name]} [{r[key]}]"
            skill_ids[(path.name, r[key])] = display
            add(path, sheet, n, r[key], "skill", {"name": display, "source_name": r[name],
                "category": r.get(category), "description": r.get(description),
                "max_level": 5 if path == FINANCE_FILE else 3})

    roles = unique(FINANCE_FILE, "Role_Master", 3, "role_id", ["role_name", "tower"])
    role_ids = set()
    for n, r in roles:
        if not r.get("role_name"):
            issue(FINANCE_FILE, "Role_Master", n, "Missing role name", r["role_id"]); continue
        role_ids.add(r["role_id"])
        add(FINANCE_FILE, "Role_Master", n, r["role_id"], "role", {
            "role_code": r["role_id"], "role_name": r["role_name"], "business_function": "Finance",
            "tower": r.get("tower"), "role_grade": r.get("role_grade"), "level_framework": "FINANCE_0_5"})
    descriptions = unique(FINANCE_FILE, "Role_Descriptions", 3, "role_id", ["role_name"])
    for n, r in descriptions:
        if r["role_id"] not in role_ids:
            issue(FINANCE_FILE, "Role_Descriptions", n, "Reference-only role description; no selected Role_Master row", r["role_id"])
        add(FINANCE_FILE, "Role_Descriptions", n, r["role_id"], "role_description", {
            "role_key": r["role_id"], "role_name": r.get("role_name"),
            "summary": r.get("role_summary_or_positioning"), "responsibilities": r.get("responsibility_extract"),
            "outputs": r.get("output_or_success_measure_extract"), "skills": r.get("skills_or_capabilities_extract"),
            "source_file": r.get("source_file"), "selection_status": r.get("selection_status")})
    rd_roles = table(RD_FILE, "Role Levels", 4, ["Band", "Readiness stage"])
    bands = [(n, r["Band"]) for n, r in rd_roles if r.get("Band") in {"B2", "B3", "B4"}]
    if len(bands) != 3 or len(set(b for _, b in bands)) != 3:
        raise ValueError("RD Role Levels: missing or duplicate bands")
    for n, band in bands:
        add(RD_FILE, "Role Levels", n, band, "role", {"role_code": f"RD-{band}",
            "role_name": band, "business_function": "DataOps", "role_grade": band,
            "level_framework": "RD_PROF_01_03"})
    levels = table(RD_FILE, "Proficiency Levels", 8, ["Level_ID", "Proficiency"])
    rd_levels = {r["Proficiency"]: int(r["Level_ID"][-2:]) for _, r in levels
                 if r.get("Level_ID") in {"PROF_01", "PROF_02", "PROF_03"}}
    if rd_levels != {"Beginner": 1, "Moderate": 2, "Expert": 3}:
        raise ValueError("RD proficiency codes require source validation")
    mappings = {}
    rows = unique(FINANCE_FILE, "Role_Skill_Map", 3, "role_skill_id",
                  ["role_id", "skill_id", "target_proficiency_level"])
    pairs = Counter((r.get("role_id"), r.get("skill_id")) for _, r in rows)
    for n, r in rows:
        key = r["role_skill_id"]
        target = r.get("target_proficiency_level")
        reason = None
        if r.get("role_id") not in role_ids: reason = "Missing role reference"
        elif (FINANCE_FILE.name, r.get("skill_id")) not in skill_ids: reason = "Missing skill reference"
        elif pairs[(r["role_id"], r["skill_id"])] != 1: reason = "Ambiguous role-skill relationship"
        elif target is not None and (type(target) is not int or target not in range(6)): reason = "Invalid target level"
        if reason:
            issue(FINANCE_FILE, "Role_Skill_Map", n, reason, r.get("skill_id")); continue
        data = {"role_key": r["role_id"], "skill_key": r["skill_id"], "target_level": target,
                "target_label": str(target) if target is not None else "Not Expected",
                "is_expected": target is not None and target > 0}
        add(FINANCE_FILE, "Role_Skill_Map", n, key, "mapping", data)
        mappings[(r["role_id"], r["skill_id"], target)] = key
    for n, r in unique(RD_FILE, "Role Skill Matrix", 4, "Skill_ID", ["B2 target", "B3 target", "B4 target"]):
        for _, band in bands:
            label = r.get(f"{band} target")
            if (RD_FILE.name, r["Skill_ID"]) not in skill_ids or (label is not None and label not in rd_levels):
                issue(RD_FILE, "Role Skill Matrix", n, "Missing skill or unsupported target", r["Skill_ID"]); continue
            add(RD_FILE, "Role Skill Matrix", n, f"{band}:{r['Skill_ID']}", "mapping", {
                "role_key": band, "skill_key": r["Skill_ID"], "target_level": rd_levels.get(label),
                "target_label": label or "Not Expected", "is_expected": label is not None})
    finance_questions = unique(FINANCE_FILE, "Assessment_QBank", 3, "question_id",
                      ["role_id", "skill_id", "question_text", "correct_option", "sme_review_status"])
    for n, r in finance_questions:
        options = {c: r.get(f"option_{c.lower()}") for c in "ABCD"}
        reason = None
        mapping = mappings.get((r.get("role_id"), r.get("skill_id"), r.get("target_proficiency_level")))
        if not r.get("question_text") or not all(isinstance(v, str) and v.strip() for v in options.values()) or r.get("correct_option") not in options:
            reason = "Incomplete question options or answer key"
        elif len(set(options.values())) != 4: reason = "Duplicate answer options"
        elif mapping is None: reason = "Missing exact role-skill-target mapping"
        elif r.get("sme_review_status") not in {"Approved", "Pending SME Review", "Needs Rewrite"}: reason = "Unsupported review status"
        if reason:
            issue(FINANCE_FILE, "Assessment_QBank", n, reason, r["question_id"]); continue
        approved = r["sme_review_status"] == "Approved" and r.get("approved_for_schedule") == "Yes"
        add(FINANCE_FILE, "Assessment_QBank", n, r["question_id"], "question", {
            "mapping_key": mapping, "question_text": r["question_text"], "options": options,
            "correct_answer": r["correct_option"], "skill_level": r["target_proficiency_level"],
            "status": "approved" if approved else "pending_review", "source_status": r["sme_review_status"],
            "is_critical": r.get("critical_flag") == "Yes", "source_label": r.get("data_label"),
            "difficulty": r.get("difficulty"), "ai_confidence": r.get("ai_confidence"),
            "explanation": "Answer key supplied by the workbook; no separate rationale supplied."})
    catalogue = unique(FINANCE_FILE, "Training_Catalogue", 3, "course_id", ["course_title"])
    courses = {r["course_id"]: (n, r) for n, r in catalogue}
    for n, r in catalogue:
        add(FINANCE_FILE, "Training_Catalogue", n, r["course_id"], "course",
            {str(k): v for k, v in r.items() if k is not None} | {"validation_status": "Pending source validation"})
    for n, r in unique(FINANCE_FILE, "Skill_Gaps_TNI", 3, "gap_id",
                       ["user_id", "role_id", "skill_id", "recommended_course_id"]):
        add(FINANCE_FILE, "Skill_Gaps_TNI", n, r["gap_id"], "tni",
            {str(k): v for k, v in r.items() if k is not None} | {"validation_status": "Pending source validation"})

    for n, r in unique(FINANCE_FILE, "Training_Skill_Map", 3, "training_skill_map_id", ["course_id", "skill_id"]):
        course = courses.get(r.get("course_id"))
        if (FINANCE_FILE.name, r.get("skill_id")) not in skill_ids or not course or not course[1].get("course_title"):
            issue(FINANCE_FILE, "Training_Skill_Map", n, "Missing skill or course reference", r["training_skill_map_id"]); continue
        add(FINANCE_FILE, "Training_Skill_Map", n, r["training_skill_map_id"], "learning", {
            "skill_key": r["skill_id"], "resource_id": f"finance:{r['training_skill_map_id']}",
            "title": course[1]["course_title"], "source_file": FINANCE_FILE.name,
            "source_sheet": "Training_Catalogue", "source_row": course[0],
            "mapping_sheet": "Training_Skill_Map", "mapping_row": n, "source_skill_id": r["skill_id"],
            "level_scope": r.get("level_group"),
            "review_status": f"Mapping confidence: {r.get('mapping_confidence')}; {course[1].get('data_label') or 'Approval not supplied'}"})
    for n, r in unique(RD_FILE, "SOP & Learning Mapping", 4, "Skill_ID", ["Reference title", "Primary URL"]):
        if (RD_FILE.name, r["Skill_ID"]) not in skill_ids or not r.get("Reference title"):
            issue(RD_FILE, "SOP & Learning Mapping", n, "Missing skill or reference title", r["Skill_ID"]); continue
        url = str(r.get("Primary URL") or "")
        available = url.startswith("https://")
        if not available:
            issue(RD_FILE, "SOP & Learning Mapping", n, "Learning reference has no URL; metadata retained as unavailable", r["Skill_ID"])
        add(RD_FILE, "SOP & Learning Mapping", n, r["Skill_ID"], "learning", {
            "skill_key": r["Skill_ID"], "resource_id": f"dataops:{r['Skill_ID']}", "title": r["Reference title"],
            # URLs remain audit metadata and are never treated as ingested knowledge.
            "url": url if available else None, "original_reference": url if available else None,
            "source_type": r.get("Source type"), "owner": r.get("Owner"),
            "source_file": RD_FILE.name, "source_sheet": "SOP & Learning Mapping",
            "source_row": n, "mapping_sheet": "SOP & Learning Mapping", "mapping_row": n,
            "source_skill_id": r["Skill_ID"], "review_status": r.get("Review status"),
            "availability_status": "Content not supplied", "ingestion_status": "Not ingested"})
    for n, r in unique(RD_FILE, "Source Register", 4, "Source_ID", ["Source", "Purpose"]):
        location = str(r.get("Location / URL") or "")
        add(RD_FILE, "Source Register", n, r["Source_ID"], "source_metadata", {
            "title": r.get("Source"), "source_type": r.get("Purpose"), "owner": r.get("Owner"),
            "original_reference": location or None, "source_status": r.get("Status"),
            "last_validated_date": str(r.get("As of")) if r.get("As of") is not None else None,
            "usage_note": r.get("Usage note"), "availability_status": "Content not supplied",
            "ingestion_status": "Not ingested"})
    workbook = load_workbook(RD_FILE, read_only=True, data_only=True)
    try:
        ws = workbook["Question Bank & Difficulty"]
        metadata = {str(row[0] or "").strip(): row[1] for row in ws.iter_rows(values_only=True) if row[0]}
    finally:
        workbook.close()
    add(RD_FILE, "Question Bank & Difficulty", 5, "Faiza Question Bank - All Items", "question_source", {
        "title": metadata.get("Authoritative source"), "url": metadata.get("Live URL"),
        "use_rule": metadata.get("Use rule"), "availability_status": "Unavailable — excluded from MVP",
        "audit_reference_only": True, "questions_imported": 0})
    for book in (FINANCE_FILE.name, RD_FILE.name):
        report["counts"][book] = dict(Counter(r["entity_type"] for r in records if r["workbook"] == book))
    report["roles"] = [{"workbook": r["workbook"], "key": r["source_key"], **r["details"]}
                       for r in records if r["entity_type"] == "role"]
    report["skills"] = [{"workbook": r["workbook"], "key": r["source_key"],
                         "name": r["details"]["source_name"], "category": r["details"]["category"]}
                        for r in records if r["entity_type"] == "skill"]
    report["question_source_statuses"] = dict(Counter(r["details"]["source_status"]
        for r in records if r["entity_type"] == "question"))
    report["finance_question_inventory"] = {
        "total_rows": len(finance_questions),
        "incomplete_rows": sum(i["sheet"] == "Assessment_QBank" and
            i["reason"] == "Incomplete question options or answer key" for i in report["issues"]),
        "safely_importable": sum(r["workbook"] == FINANCE_FILE.name and
            r["entity_type"] == "question" for r in records),
    }
    report["issue_counts"] = dict(Counter(i["reason"] for i in report["issues"]))
    report["expected_targets"] = dict(Counter(r["workbook"] for r in records
        if r["entity_type"] == "mapping" and r["details"]["is_expected"]))
    report["duplicate_skill_names"] = {name: count for name, count in names.items() if count > 1}
    report["missing_finance_skill_ids"] = sorted({i["key"] for i in report["issues"]
        if i["workbook"] == FINANCE_FILE.name and i["reason"] == "Missing skill reference"})
    report["limitations"] = ["Faiza Microsoft List: Unavailable — excluded from MVP; retained as an audit reference only; imported records: 0.",
        "Score-to-proficiency policy unavailable in both workbooks. Synthetic example results are not a scoring policy.",
        "Learning links and Finance prototype mappings are not proof of approved SOP content or course suitability.",
        "Starting size is 20 with 6 Difficult, 8 Moderate and 6 Easy questions, but no Finance role has 20 approved questions and source difficulty uses Role Ready/Advanced.",
        "Twenty questions cannot cover every required B3/B4 skill; the approved allocation rule is unresolved.",
        "Reassessment requires at least 70% new questions, but the approved source pool is insufficient.",
        "Criticality versus difficulty is not formally defined; no automatic equivalence is applied.",
        "No approved AI-confidence scale or threshold is supplied; confidence is review triage only.",
        "The subjective-scoring repository and MyAcademy document content are not supplied.",
        "The more-than-90-percent document-match comments do not define a semantic, rubric or exact match method.",
        "The combined Beginner/Moderate/Expert formula is incomplete, including the Moderate level and how objective, behavioural and evidence results combine."]
    report["assessment_rules"] = {"starting_question_count": 20,
        "difficulty_distribution": {"Difficult": 6, "Moderate": 8, "Easy": 6},
        "critical_incorrect_action": "insight_and_review_only", "minimum_reassessment_new_percent": 70,
        "ai_confidence_use": "review_triage_only", "expert_requires": "complex-case evidence plus manager and SME confirmation"}
    return records, report


def import_sources(db, actor_id=None):
    records, report = source_plan()
    entities = {}
    changed = 0
    for data in records:
        identity = {k: data[k] for k in ("workbook", "sheet", "source_key")}
        old = db.query(models.SourceRecord).filter_by(**identity).first()
        fingerprint = hashlib.sha256(json.dumps([data["source_row"], data["details"]], sort_keys=True, default=str).encode()).hexdigest()
        kind, values = data["entity_type"], data["details"]
        model = {"role": models.Role, "skill": models.Skill, "mapping": models.RoleSkillMap, "question": models.Question}.get(kind)
        row = db.get(model, old.entity_id) if old and model else None
        if old and old.fingerprint == fingerprint:
            entities[(data["workbook"], kind, data["source_key"])] = row
            continue
        if model:
            if row is None:
                row = model(); db.add(row)
            if kind == "role":
                for k, v in values.items(): setattr(row, k, v)
            elif kind == "skill":
                for k in ("name", "description", "max_level"): setattr(row, k, values[k])
            elif kind == "mapping":
                row.role_id = entities[(data["workbook"], "role", values["role_key"])].id
                row.skill_id = entities[(data["workbook"], "skill", values["skill_key"])].id
                for k in ("target_level", "target_label", "is_expected"): setattr(row, k, values[k])
            elif kind == "question":
                mapping = entities[(data["workbook"], "mapping", values["mapping_key"])]
                row.skill_id = mapping.skill_id
                for k in ("question_text", "options", "correct_answer", "skill_level", "status", "explanation"): setattr(row, k, values[k])
                row.rag_source = (f"{data['workbook']} / {data['sheet']} / row {data['source_row']} / "
                    f"{data['source_key']}; {values['source_label']}; difficulty={values['difficulty']}; "
                    f"AI confidence={values['ai_confidence']} (review triage only)")
                row.review_comment = f"Workbook status: {values['source_status']}"
            db.flush()
            if kind == "question":
                scope = db.get(models.QuestionScope, row.id)
                if scope is None: db.add(models.QuestionScope(question_id=row.id, role_skill_map_id=mapping.id))
                else: scope.role_skill_map_id = mapping.id
                record_question(db, row, actor_id, "source_imported", is_critical=values["is_critical"])
        elif kind == "role_description":
            selected_role = entities.get((data["workbook"], "role", values["role_key"]))
            row = selected_role
        if old is None:
            old = models.SourceRecord(**data, fingerprint=fingerprint, entity_id=row.id if row else None)
            db.add(old)
        else:
            old.details, old.fingerprint, old.source_row = values, fingerprint, data["source_row"]
        db.flush()
        entities[(data["workbook"], kind, data["source_key"])] = row
        changed += 1
    # Source disappearance is never silently deleted or reused for new assignments.
    active = {(r['workbook'], r['sheet'], r['source_key']) for r in records}
    for old in db.query(models.SourceRecord):
        if (old.workbook, old.sheet, old.source_key) not in active:
            report["issues"].append({"workbook": old.workbook, "sheet": old.sheet, "key": old.source_key,
                                     "reason": "Previously imported row now unavailable; requires reconciliation"})
            if old.entity_type == "mapping": db.get(models.RoleSkillMap, old.entity_id).is_expected = False
            if old.entity_type == "question": db.get(models.Question, old.entity_id).status = "pending_review"
            old.fingerprint = ""
    if changed:
        audit(db, actor_id, "sources.imported", "source", None, details={"changed": changed, "counts": report["counts"]})
    report["changed"] = changed
    return report


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if args.apply:
        import config
        from database import Base, engine, SessionLocal
        Base.metadata.create_all(engine)
        with SessionLocal.begin() as db:
            result = import_sources(db)
    else:
        _, result = source_plan()
    print(json.dumps(result, indent=2))
