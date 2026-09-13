import json
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

import config
from offline_rag import (NO_CONTEXT, ManifestEntry, RagError, chunk_segments,
                         extract_document, ingest_manifest, public_source_status, retrieve, validate_entry)


def _pdf(path: Path, text="Synthetic PDF fixture"):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 5 0 R >> >> /MediaBox [0 0 612 792] /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    data = bytearray(b"%PDF-1.4\n"); offsets = [0]
    for i, obj in enumerate(objects, 1): offsets.append(len(data)); data += f"{i} 0 obj\n".encode()+obj+b"\nendobj\n"
    xref=len(data); data += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    data += b"".join(f"{o:010} 00000 n \n".encode() for o in offsets[1:])
    data += f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode(); path.write_bytes(data)


def test_pdf_docx_pptx_and_text_extraction(tmp_path):
    pdf=tmp_path/"a.pdf"; _pdf(pdf); assert "Synthetic PDF" in extract_document(pdf)[0]["text"]
    docx=tmp_path/"a.docx"; d=Document(); d.add_paragraph("Synthetic DOCX fixture"); d.save(docx); assert "DOCX" in extract_document(docx)[0]["text"]
    pptx=tmp_path/"a.pptx"; p=Presentation(); s=p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text="Synthetic PPTX fixture"; p.save(pptx); assert "PPTX" in extract_document(pptx)[0]["text"]
    txt=tmp_path/"a.txt"; txt.write_text("Synthetic text",encoding="utf8"); assert extract_document(txt)[0]["text"] == "Synthetic text"


def test_chunking_is_deterministic_and_ids_are_stable():
    segments=[{"reference":"page 1","text":"word "*100}]
    one=chunk_segments(segments,"SRC","abc",max_chars=100)
    two=chunk_segments(segments,"SRC","abc",max_chars=100)
    assert one == two and len(one)>1 and all(c["chunk_id"].startswith("CHK_") for c in one)


def test_validation_guards(tmp_path, monkeypatch):
    monkeypatch.setattr(config,"RAG_MAX_FILE_BYTES",1000); monkeypatch.setattr(config,"LLM_PROVIDER","mock"); monkeypatch.setattr(config,"ALLOW_SYNTHETIC_RAG",True)
    def entry(**kw):
        values=dict(source_id="SYNTHETIC_X",filename="a.txt",skill_ids=["RD_ADM_01"],source_type="test",approved_for_generation=True,synthetic_only=True,title="Fixture — Test Fixture Only")
        values.update(kw); return ManifestEntry(**values)
    with pytest.raises(RagError,match="missing"): validate_entry(entry(),tmp_path,{}, {"RD_ADM_01"})
    (tmp_path/"a.txt").write_text("x"); assert validate_entry(entry(),tmp_path,{}, {"RD_ADM_01"}).name == "a.txt"
    with pytest.raises(RagError,match="traversal"): validate_entry(entry(filename="../a.txt"),tmp_path,{}, {"RD_ADM_01"})
    with pytest.raises(RagError,match="Unknown Skill_ID"): validate_entry(entry(skill_ids=["BAD"]),tmp_path,{}, {"RD_ADM_01"})
    bad=tmp_path/"a.exe"; bad.write_text("x")
    with pytest.raises(RagError,match="Unsupported"): validate_entry(entry(filename="a.exe"),tmp_path,{}, {"RD_ADM_01"})


def test_empty_and_broken_documents_fail(tmp_path):
    empty=tmp_path/"empty.txt"; empty.write_text("");
    with pytest.raises(RagError,match="empty"): extract_document(empty)
    broken=tmp_path/"broken.pdf"; broken.write_text("bad")
    with pytest.raises(RagError,match="extraction failed"): extract_document(broken)


def test_ingestion_duplicate_status_and_public_metadata(tmp_path, monkeypatch):
    monkeypatch.setattr(config,"LLM_PROVIDER","mock"); monkeypatch.setattr(config,"ALLOW_SYNTHETIC_RAG",True)
    (tmp_path/"a.txt").write_text("same synthetic content"); (tmp_path/"b.txt").write_text("same synthetic content")
    base={"skill_ids":["RD_ADM_01"],"source_type":"test","approved_for_generation":True,"synthetic_only":True,"function":"DataOps","title":"Fixture — Test Fixture Only"}
    (tmp_path/"manifest.json").write_text(json.dumps({"sources":[{"source_id":"SYNTHETIC_A","filename":"a.txt",**base},{"source_id":"SYNTHETIC_B","filename":"b.txt",**base}]}))
    result=ingest_manifest(source_dir=tmp_path,index_dir=tmp_path/"indexindex")
    assert result["sources"][0]["ingestion_status"] == "Ingested"
    assert "Duplicate content" in result["sources"][1]["limitation"]
    assert "text" not in json.dumps(public_source_status(result))


def test_retrieval_is_skill_and_function_scoped(monkeypatch):
    monkeypatch.setattr(config,"LLM_PROVIDER","mock"); monkeypatch.setattr(config,"ALLOW_SYNTHETIC_RAG",True)
    chunk={"chunk_id":"CHK_1","text":"unique validation record","source_id":"S","source_title":"T","skill_ids":["RD_ADM_01"],"document_filename":"a.txt","synthetic_only":True,"function":"DataOps","approved_for_generation":True,"ingestion_status":"Ingested","intended_proficiencies":["Moderate"]}
    result=retrieve(function="DataOps",skill_id="RD_ADM_01",intended_proficiency="Moderate",query="validation",index={"chunks":[chunk]})
    assert result["chunk_references"] == ["CHK_1"] and result["selected_chunks"][0]["retrieval_score"] > 0
    for values in [("Finance","RD_ADM_01"),("DataOps","RD_OTHER")]:
        with pytest.raises(RagError,match=NO_CONTEXT): retrieve(function=values[0],skill_id=values[1],intended_proficiency="Moderate",query="validation",index={"chunks":[chunk]})


def test_synthetic_rejected_outside_mock(monkeypatch,tmp_path):
    (tmp_path/"a.txt").write_text("x")
    entry=ManifestEntry(source_id="SYNTHETIC_X",filename="a.txt",skill_ids=["RD_ADM_01"],source_type="test",approved_for_generation=True,synthetic_only=True,title="Fixture — Test Fixture Only")
    monkeypatch.setattr(config,"LLM_PROVIDER","luna"); monkeypatch.setattr(config,"ALLOW_SYNTHETIC_RAG",True)
    with pytest.raises(RagError,match="disabled"): validate_entry(entry,tmp_path,{}, {"RD_ADM_01"})


def test_dataops_generation_is_grounded_pending_and_not_auto_approved(client, monkeypatch):
    from database import SessionLocal
    import models
    role=client.post("/roles",json={"role_code":"RD-B2-X","role_name":"B2 Test","business_function":"DataOps","role_grade":"B2"}).json()
    skill=client.post("/skills",json={"name":"Synthetic RD Skill","description":"fixture","max_level":3}).json()
    mapping=client.post("/role-skill-maps",json={"role_id":role["id"],"skill_id":skill["id"],"target_level":2,"target_label":"Moderate"}).json()
    with SessionLocal.begin() as db:
        db.add(models.SourceRecord(workbook="overall_rd.xlsx",sheet="Skill Master",source_key="RD_ADM_01",entity_type="skill",entity_id=skill["id"],source_row=5,fingerprint="x",details={"name":"Synthetic RD Skill"}))
        db.add(models.SourceRecord(workbook="overall_rd.xlsx",sheet="Role Skill Matrix",source_key="B2:RD_ADM_01",entity_type="mapping",entity_id=mapping["id"],source_row=5,fingerprint="y",details={"is_expected":True}))
    grounding={"selected_chunks":[{"chunk_id":"CHK_1","text":"Synthetic rule requires a unique practice identifier.","source_id":"SYNTHETIC_A","reference":"section 1","document_filename":"a.txt","synthetic_only":True}],"source_ids":["SYNTHETIC_A"],"document_references":["a.txt"],"chunk_references":["CHK_1"],"limitations":[]}
    monkeypatch.setattr("question_service.retrieve",lambda **kwargs: grounding)
    response=client.post("/questions/generate",json={"role_skill_map_id":mapping["id"],"question_count":1,"difficulty":"Moderate"})
    assert response.status_code == 201, response.text
    question=response.json()[0]
    assert question["status"] == "pending_review" and question["source_skill_id"] == "RD_ADM_01"
    assert question["chunk_references"] == ["CHK_1"] and question["ai_confidence"] == "mock-deterministic"
    with SessionLocal() as db:
        assert db.query(models.Question).filter_by(id=question["id"],status="approved").first() is None
