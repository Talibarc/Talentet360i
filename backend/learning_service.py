"""Exact source joins only. Missing/ambiguous mappings never become invented courses."""

from pathlib import Path
from zipfile import BadZipFile
from openpyxl import load_workbook

from excel_loader import FINANCE_FILE, RD_FILE
from schemas import LearningResource


def _rows(path: Path, sheet: str, header_row: int) -> list[tuple[int, dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = workbook[sheet]
        headers = [str(c.value).strip() if c.value is not None else None for c in ws[header_row]]
        return [(index, {key: value for key, value in zip(headers, values) if key})
                for index, values in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                               start=header_row + 1)
                if any(value is not None for value in values)]
    finally:
        workbook.close()


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def _matches(value, name: str) -> bool:
    return _text(value).casefold() == name.strip().casefold()


def get_learning_resources(skill_name: str, business_function: str) -> tuple[list[LearningResource], str]:
    """Return skill-mapped candidates, not an inferred level-specific curriculum."""
    resources = []
    try:
        if business_function.strip().casefold() == "finance":
            masters = _rows(FINANCE_FILE, "Skill_Master", 3)
            ids = {_text(r.get("skill_id")) for _, r in masters
                   if _matches(r.get("skill"), skill_name) and r.get("skill_id")}
            if len(ids) != 1:
                return [], "No unique source skill mapping available"
            skill_id = next(iter(ids))
            catalogue = _rows(FINANCE_FILE, "Training_Catalogue", 3)
            for mapping_row, mapping in _rows(FINANCE_FILE, "Training_Skill_Map", 3):
                if _text(mapping.get("skill_id")) != skill_id:
                    continue
                course_id = _text(mapping.get("course_id"))
                courses = [(row, r) for row, r in catalogue
                           if course_id and _text(r.get("course_id")) == course_id]
                if len(courses) != 1 or not _text(courses[0][1].get("course_title")):
                    continue
                row, course = courses[0]
                resources.append(LearningResource(
                    resource_id=f"finance:{course_id}:{mapping_row}", title=_text(course["course_title"]),
                    source_file=FINANCE_FILE.name, source_sheet="Training_Catalogue", source_row=row,
                    mapping_sheet="Training_Skill_Map", mapping_row=mapping_row, source_skill_id=skill_id,
                    level_scope=_text(mapping.get("level_group")) or None,
                    review_status=_text(mapping.get("mapping_confidence")) or None,
                ))
        elif business_function.strip().casefold() == "dataops":
            masters = _rows(RD_FILE, "Skill Master", 4)
            ids = {_text(r.get("Skill_ID")) for _, r in masters
                   if _matches(r.get("Skill"), skill_name) and r.get("Skill_ID")}
            if len(ids) != 1:
                return [], "No unique source skill mapping available"
            skill_id = next(iter(ids))
            for row, mapping in _rows(RD_FILE, "SOP & Learning Mapping", 4):
                if _text(mapping.get("Skill_ID")) != skill_id or not _text(mapping.get("Reference title")):
                    continue
                url = _text(mapping.get("Primary URL"))
                resources.append(LearningResource(
                    resource_id=f"dataops:{skill_id}:{row}", title=_text(mapping["Reference title"]),
                    url=url if url.startswith(("https://", "http://")) else None,
                    source_file=RD_FILE.name, source_sheet="SOP & Learning Mapping", source_row=row,
                    mapping_sheet="SOP & Learning Mapping", mapping_row=row, source_skill_id=skill_id,
                    review_status=_text(mapping.get("Review status")) or None,
                ))
        else:
            return [], "No workbook configured for this business function"
    except (OSError, KeyError, ValueError, BadZipFile):
        return [], "Learning workbook or required sheet unavailable"
    return resources, ("Skill-mapped resources; level suitability and approval require review"
                       if resources else "No supported learning mapping available")
