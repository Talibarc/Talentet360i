from pathlib import Path
from typing import Any
 
from openpyxl import load_workbook
 
 
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
 
FINANCE_FILE = DATA_DIR / "finance_assessment.xlsx"
RD_FILE = DATA_DIR / "overall_rd.xlsx"
 
 
def read_sheet(
    file_path: Path,
    sheet_name: str,
    header_row: int,
) -> list[dict[str, Any]]:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path.name}")
 
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )
 
    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet '{sheet_name}' not found in {file_path.name}"
        )
 
    worksheet = workbook[sheet_name]
 
    headers = [
        str(cell.value).strip() if cell.value is not None else None
        for cell in worksheet[header_row]
    ]
 
    records = []
 
    for values in worksheet.iter_rows(
        min_row=header_row + 1,
        values_only=True,
    ):
        if not any(value is not None for value in values):
            continue
 
        record = {
            header: value
            for header, value in zip(headers, values)
            if header is not None
        }
        records.append(record)
 
    workbook.close()
    return records
 
 
def validate_workbooks() -> dict[str, Any]:
    finance_roles = read_sheet(
        FINANCE_FILE,
        "Role_Master",
        3,
    )
    finance_skills = read_sheet(
        FINANCE_FILE,
        "Skill_Master",
        3,
    )
    finance_mappings = read_sheet(
        FINANCE_FILE,
        "Role_Skill_Map",
        3,
    )
 
    rd_skills = read_sheet(
        RD_FILE,
        "Skill Master",
        4,
    )
    rd_matrix = read_sheet(
        RD_FILE,
        "Role Skill Matrix",
        4,
    )
    rd_learning = read_sheet(
        RD_FILE,
        "SOP & Learning Mapping",
        4,
    )
 
    finance_skill_ids = {
        row["skill_id"]
        for row in finance_skills
        if row.get("skill_id")
    }
 
    missing_skill_ids = sorted({
        row["skill_id"]
        for row in finance_mappings
        if (
            row.get("skill_id")
            and row["skill_id"] not in finance_skill_ids
        )
    })
 
    issues = []
 
    if missing_skill_ids:
        issues.append({
            "source": "finance_assessment",
            "issue": "Role mappings reference missing skills",
            "missing_skill_count": len(missing_skill_ids),
            "missing_skill_ids": missing_skill_ids,
        })
 
    return {
        "status": "needs_review" if issues else "ready",
        "finance_assessment": {
            "roles": len(finance_roles),
            "skills": len(finance_skills),
            "role_skill_mappings": len(finance_mappings),
        },
        "overall_rd": {
            "skills": len(rd_skills),
            "role_skill_mappings": len(rd_matrix),
            "learning_mappings": len(rd_learning),
        },
        "issues": issues,
    }