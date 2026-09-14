"""Offline, manifest-bound document ingestion and deterministic retrieval.

This module never opens URLs. Files are resolved beneath RAG_SOURCE_DIR and the
generated JSON index is written beneath RAG_INDEX_DIR, which is excluded from Git.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

from docx import Document
from pydantic import BaseModel, ConfigDict, Field, model_validator
from pypdf import PdfReader
from pptx import Presentation

import config


NO_CONTEXT = "No approved source content available for this skill"
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt"}
INDEX_FILENAME = "index.json"


class RagError(ValueError):
    """A safe validation, ingestion, or retrieval failure."""


class ManifestEntry(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)

    source_id: str = Field(min_length=1, max_length=200)
    filename: str = Field(min_length=1, max_length=500)
    skill_ids: list[str] = Field(min_length=1)
    source_type: str = Field(min_length=1, max_length=100)
    approved_for_generation: bool
    synthetic_only: bool = False
    function: Literal["RD", "DataOps"] = "DataOps"
    title: str | None = Field(default=None, max_length=500)
    owner: str | None = Field(default=None, max_length=300)
    original_reference: str | None = Field(default=None, max_length=2000)
    role_bands: list[str] = Field(default_factory=list)
    intended_proficiencies: list[str] = Field(default_factory=list)
    last_validated_date: str | None = None

    @model_validator(mode="after")
    def validate_lists(self):
        self.skill_ids = sorted(set(self.skill_ids))
        self.role_bands = sorted(set(self.role_bands))
        self.intended_proficiencies = sorted(set(self.intended_proficiencies))
        if self.synthetic_only and not self.source_id.startswith("SYNTHETIC_"):
            raise ValueError("Synthetic source IDs must start with SYNTHETIC_")
        if self.synthetic_only and "Test Fixture Only" not in (self.title or ""):
            raise ValueError("Synthetic source titles must contain 'Test Fixture Only'")
        return self


class Manifest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    sources: list[ManifestEntry] = Field(min_length=1)

    @model_validator(mode="after")
    def unique_entries(self):
        ids = [entry.source_id for entry in self.sources]
        files = [entry.filename.casefold() for entry in self.sources]
        if len(ids) != len(set(ids)):
            raise ValueError("Duplicate source ID in manifest")
        if len(files) != len(set(files)):
            raise ValueError("Duplicate filename in manifest")
        return self


def _inside(root: Path, relative: str, *, label: str) -> Path:
    candidate = Path(relative)
    if candidate.is_absolute():
        raise RagError(f"{label} must be relative to RAG_SOURCE_DIR")
    root = root.resolve()
    resolved = (root / candidate).resolve()
    if not resolved.is_relative_to(root):
        raise RagError(f"{label} path traversal rejected")
    return resolved


def _text(value: str | None) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def extract_document(path: Path) -> list[dict[str, str]]:
    """Extract ordered text segments with the best available local reference."""
    extension = path.suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise RagError(f"Unsupported file extension: {extension or '(none)'}")
    try:
        if extension == ".pdf":
            segments = [
                {"reference": f"page {number}", "text": _text(page.extract_text())}
                for number, page in enumerate(PdfReader(str(path)).pages, 1)
            ]
        elif extension == ".docx":
            document = Document(str(path))
            segments = [
                {"reference": f"paragraph {number}", "text": _text(paragraph.text)}
                for number, paragraph in enumerate(document.paragraphs, 1)
            ]
            for table_number, table in enumerate(document.tables, 1):
                for row_number, row in enumerate(table.rows, 1):
                    segments.append({
                        "reference": f"table {table_number}, row {row_number}",
                        "text": _text(" | ".join(cell.text for cell in row.cells)),
                    })
        elif extension == ".pptx":
            presentation = Presentation(str(path))
            segments = []
            for slide_number, slide in enumerate(presentation.slides, 1):
                content = [shape.text for shape in slide.shapes if hasattr(shape, "text")]
                segments.append({"reference": f"slide {slide_number}", "text": _text(" ".join(content))})
        else:
            value = path.read_text(encoding="utf-8-sig")
            segments = [
                {"reference": f"section {number}", "text": _text(part)}
                for number, part in enumerate(re.split(r"\n\s*\n", value), 1)
            ]
    except RagError:
        raise
    except Exception as error:
        raise RagError(f"Document extraction failed: {type(error).__name__}") from None
    segments = [segment for segment in segments if segment["text"]]
    if not segments:
        raise RagError("Document is empty")
    return segments


def chunk_segments(
    segments: list[dict[str, str]], source_id: str, file_hash: str, *, max_chars: int = 1000
) -> list[dict[str, Any]]:
    if max_chars < 100:
        raise RagError("Chunk size must be at least 100 characters")
    chunks: list[dict[str, Any]] = []
    sequence = 0
    for segment in segments:
        words = segment["text"].split()
        current: list[str] = []
        current_size = 0
        pieces: list[str] = []
        for word in words:
            addition = len(word) + (1 if current else 0)
            if current and current_size + addition > max_chars:
                pieces.append(" ".join(current))
                current, current_size = [], 0
            current.append(word)
            current_size += len(word) + (1 if len(current) > 1 else 0)
        if current:
            pieces.append(" ".join(current))
        for text in pieces:
            sequence += 1
            identity = json.dumps(
                [source_id, file_hash, sequence, segment["reference"], text],
                ensure_ascii=False,
                separators=(",", ":"),
            )
            chunks.append({
                "chunk_id": "CHK_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
                "chunk_sequence": sequence,
                "reference": segment["reference"],
                "text": text,
            })
    return chunks


def workbook_registry() -> tuple[dict[str, dict[str, Any]], set[str]]:
    """Return stable RD source metadata and source Skill_IDs without opening URLs."""
    from source_import import source_plan

    records, _ = source_plan()
    source_records: dict[str, dict[str, Any]] = {}
    skills: set[str] = set()
    for record in records:
        if record["workbook"] != "overall_rd.xlsx":
            continue
        if record["entity_type"] == "skill":
            skills.add(record["source_key"])
        if record["entity_type"] == "source_metadata":
            source_records[record["source_key"]] = record["details"]
        if record["entity_type"] == "learning":
            source_records[record["details"]["resource_id"]] = {
                "title": record["details"]["title"],
                "source_type": record["details"].get("source_type"),
                "owner": record["details"].get("owner"),
                "original_reference": record["details"].get("original_reference"),
                "availability_status": "Content not supplied",
                "last_validated_date": None,
            }
    return source_records, skills


def validate_entry(entry: ManifestEntry, root: Path, registry: dict[str, dict], skill_ids: set[str]) -> Path:
    path = _inside(root, entry.filename, label="Source file")
    if path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
        raise RagError(f"Unsupported file extension: {path.suffix or '(none)'}")
    if not path.is_file():
        raise RagError("Source file is missing")
    if path.stat().st_size > config.RAG_MAX_FILE_BYTES:
        raise RagError("Source file exceeds configured size limit")
    if not entry.approved_for_generation:
        raise RagError("Source is not approved for generation")
    unknown_skills = sorted(set(entry.skill_ids) - skill_ids)
    if unknown_skills:
        raise RagError(f"Unknown Skill_ID: {', '.join(unknown_skills)}")
    if entry.synthetic_only:
        if config.LLM_PROVIDER != "mock" or not config.ALLOW_SYNTHETIC_RAG:
            raise RagError("Synthetic source is disabled for this environment")
    elif entry.source_id not in registry:
        raise RagError("Source ID is not present in the imported RD source registry")
    return path


def _load_manifest(path: Path, root: Path) -> Manifest:
    resolved = _inside(root, str(path), label="Manifest") if not path.is_absolute() else path.resolve()
    if path.is_absolute() and not resolved.is_relative_to(root.resolve()):
        raise RagError("Manifest must be stored under RAG_SOURCE_DIR")
    try:
        return Manifest.model_validate_json(resolved.read_text(encoding="utf-8"))
    except RagError:
        raise
    except Exception as error:
        raise RagError(f"Manifest validation failed: {type(error).__name__}") from None


def ingest_manifest(
    manifest_path: str | Path = "manifest.json",
    *,
    source_dir: Path | None = None,
    index_dir: Path | None = None,
) -> dict[str, Any]:
    root = (source_dir or config.RAG_SOURCE_DIR).resolve()
    destination = (index_dir or config.RAG_INDEX_DIR).resolve()
    manifest = _load_manifest(Path(manifest_path), root)
    registry, skill_ids = workbook_registry()
    sources: list[dict[str, Any]] = []
    chunks: list[dict[str, Any]] = []
    hashes: dict[str, str] = {}
    timestamp = datetime.now(timezone.utc).isoformat()

    for entry in manifest.sources:
        base = {
            **entry.model_dump(),
            "function": "DataOps",
            "availability_status": "Content not supplied",
            "ingestion_status": "Ingestion failed",
            "chunk_count": 0,
            "content_hash": None,
            "version": None,
            "ingested_at": timestamp,
            "limitation": None,
        }
        try:
            path = validate_entry(entry, root, registry, skill_ids)
            file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
            base.update(availability_status="Available locally", content_hash=file_hash, version=file_hash[:12])
            if file_hash in hashes:
                raise RagError(f"Duplicate content matches source {hashes[file_hash]}")
            segments = extract_document(path)
            document_chunks = chunk_segments(segments, entry.source_id, file_hash)
            metadata = registry.get(entry.source_id, {})
            for chunk in document_chunks:
                chunk.update(
                    source_id=entry.source_id,
                    source_title=entry.title or metadata.get("title") or entry.source_id,
                    skill_ids=entry.skill_ids,
                    document_filename=entry.filename,
                    file_content_hash=file_hash,
                    version=file_hash[:12],
                    synthetic_only=entry.synthetic_only,
                    function="DataOps",
                    role_bands=entry.role_bands,
                    intended_proficiencies=entry.intended_proficiencies,
                    approved_for_generation=True,
                    ingestion_status="Ingested",
                    ingested_at=timestamp,
                )
            chunks.extend(document_chunks)
            hashes[file_hash] = entry.source_id
            base.update(availability_status="Ingested", ingestion_status="Ingested", chunk_count=len(document_chunks))
        except RagError as error:
            base["limitation"] = str(error)
        sources.append(base)

    payload = {
        "version": 1,
        "created_at": timestamp,
        "source_dir": "Configured local source directory",
        "sources": sources,
        "chunks": chunks,
    }
    destination.mkdir(parents=True, exist_ok=True)
    temporary = destination / (INDEX_FILENAME + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    temporary.replace(destination / INDEX_FILENAME)
    return payload


def load_index(index_dir: Path | None = None) -> dict[str, Any]:
    path = (index_dir or config.RAG_INDEX_DIR) / INDEX_FILENAME
    if not path.is_file():
        return {"version": 1, "sources": [], "chunks": [], "limitation": NO_CONTEXT}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        from database import SessionLocal
        from models import SourceSkillOverride
        with SessionLocal() as db:
            overrides = {row.source_id: row.skill_ids for row in db.query(SourceSkillOverride)}
        for row in payload.get("sources", []) + payload.get("chunks", []):
            if row.get("source_id") in overrides:
                row["skill_ids"] = overrides[row["source_id"]]
        return payload
    except (OSError, json.JSONDecodeError):
        return {"version": 1, "sources": [], "chunks": [], "limitation": "Local RAG index is unavailable"}


def _tokens(value: str) -> set[str]:
    return {token for token in re.findall(r"[a-z0-9]+", value.casefold()) if len(token) > 1}


def retrieve(
    *,
    function: str,
    skill_id: str,
    intended_proficiency: str | None,
    query: str,
    limit: int = 5,
    index: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if function not in {"RD", "DataOps"}:
        raise RagError(NO_CONTEXT)
    allow_synthetic = config.LLM_PROVIDER == "mock" and config.ALLOW_SYNTHETIC_RAG
    query_tokens = _tokens(query)
    eligible = []
    for chunk in (index or load_index()).get("chunks", []):
        levels = {str(value).casefold() for value in chunk.get("intended_proficiencies", [])}
        if (
            chunk.get("function") != "DataOps"
            or skill_id not in chunk.get("skill_ids", [])
            or chunk.get("ingestion_status") != "Ingested"
            or not chunk.get("approved_for_generation")
            or (chunk.get("synthetic_only") and not allow_synthetic)
            or (intended_proficiency and levels and intended_proficiency.casefold() not in levels)
        ):
            continue
        chunk_tokens = _tokens(chunk.get("text", ""))
        overlap = len(query_tokens & chunk_tokens)
        score = overlap / max(len(query_tokens), 1)
        eligible.append((score, chunk["chunk_id"], chunk))
    if not eligible:
        raise RagError(NO_CONTEXT)
    selected = [item[2] | {"retrieval_score": round(item[0], 6)}
                for item in sorted(eligible, key=lambda item: (-item[0], item[1]))[:limit]]
    return {
        "selected_chunks": selected,
        "source_ids": sorted({item["source_id"] for item in selected}),
        "document_references": sorted({item["document_filename"] for item in selected}),
        "chunk_references": [item["chunk_id"] for item in selected],
        "limitations": ["Deterministic offline lexical retrieval; no external embedding or LLM service used."],
    }


def public_source_status(index: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    """Return source metadata only; never expose extracted content."""
    result = []
    for source in (index or load_index()).get("sources", []):
        if config.LLM_PROVIDER == "luna" and source.get("synthetic_only"):
            continue
        result.append({key: source.get(key) for key in (
            "source_id", "title", "skill_ids", "source_type", "owner", "original_reference",
            "role_bands", "intended_proficiencies", "availability_status", "ingestion_status",
            "chunk_count", "content_hash", "version", "original_filename", "internal_storage_name",
            "last_validated_date", "synthetic_only", "limitation",
        )})
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Ingest approved local documents into the offline RAG index")
    parser.add_argument("--manifest", default="manifest.json", help="Manifest path relative to RAG_SOURCE_DIR")
    arguments = parser.parse_args()
    print(json.dumps(ingest_manifest(arguments.manifest), indent=2, ensure_ascii=False))
